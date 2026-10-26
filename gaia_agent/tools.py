"""Tools expose evidence and bounded computations; there is no shell/code-execution tool."""

import ast
import io
import json
import math
import operator
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from decimal import Decimal
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlsplit

from bs4 import BeautifulSoup
from ddgs import DDGS
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import InferenceAccountError, safe_error
from .net import fetch_url


def calculator(expression: str):
    if len(expression) > 2000:
        raise ValueError("Expression is too long.")
    tree = ast.parse(expression, mode="eval")
    if len(list(ast.walk(tree))) > 300:
        raise ValueError("Expression is too complex.")
    ops = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.FloorDiv: operator.floordiv,
        ast.Mod: operator.mod,
        ast.Pow: operator.pow,
    }
    funcs = {
        "sum": sum,
        "min": min,
        "max": max,
        "abs": abs,
        "round": round,
        "sqrt": math.sqrt,
        "floor": math.floor,
        "ceil": math.ceil,
        "len": len,
    }

    def visit(node):
        if isinstance(node, ast.Constant) and type(node.value) in {int, float}:
            result = node.value
        elif isinstance(node, (ast.List, ast.Tuple)):
            return [visit(n) for n in node.elts]
        elif isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            result = (-1 if isinstance(node.op, ast.USub) else 1) * visit(node.operand)
        elif isinstance(node, ast.BinOp) and type(node.op) in ops:
            left, right = visit(node.left), visit(node.right)
            if type(left) not in {int, float} or type(right) not in {int, float}:
                raise ValueError("Arithmetic operands must be numbers.")
            if isinstance(node.op, ast.Pow) and abs(right) > 1000:
                raise ValueError("Exponent is too large.")
            result = ops[type(node.op)](left, right)
        elif (
            isinstance(node, ast.Call)
            and isinstance(node.func, ast.Name)
            and node.func.id in funcs
            and not node.keywords
        ):
            result = funcs[node.func.id](*[visit(n) for n in node.args])
        else:
            raise ValueError("Only numeric arithmetic, lists, and approved math functions are allowed.")
        if not isinstance(result, (int, float)) or not math.isfinite(result) or abs(result) > 1e100:
            raise ValueError("Result is outside the supported numeric range.")
        return result

    return visit(tree.body)


def check_archive(data: bytes):
    if zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded document exceeds the size limit.")


def extract_text(data: bytes, suffix: str, formulas=False) -> str:
    check_archive(data)
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(
            f"[Page {i + 1}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages[:100])
        )
    if suffix == ".xlsx":
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=not formulas)
        parts = []
        try:
            for sheet in book:
                parts.append(f"[Sheet: {sheet.title}; cell values, 1-based row/column coordinates]")
                for row in sheet.iter_rows(
                    max_row=min(sheet.max_row, 2000), max_col=min(sheet.max_column, 80)
                ):
                    cells = [f"{c.coordinate}={c.value}" for c in row if c.value is not None]
                    if cells:
                        parts.append(" | ".join(cells))
            parts.append(
                "Formula cells use stored Excel values. Missing cached values require checking the source formulas."
            )
        finally:
            book.close()
        return "\n".join(parts)
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(
            [p.text for p in doc.paragraphs]
            + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        )
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(io.BytesIO(data))
        return "\n".join(
            f"[Slide {i + 1}] " + "\n".join(s.text for s in slide.shapes if s.has_text_frame)
            for i, slide in enumerate(deck.slides)
        )
    if suffix in {".txt", ".py", ".csv", ".json", ".md", ".html", ".xml"}:
        return data.decode("utf-8-sig", errors="replace")
    raise ValueError(
        "Use inspect_image for images, transcribe_audio for audio; this file type has no text extractor."
    )


def youtube_id(url: str) -> str:
    parsed = urlsplit(url)
    if parsed.scheme != "https" or parsed.hostname not in {
        "youtube.com",
        "www.youtube.com",
        "m.youtube.com",
        "youtu.be",
    }:
        raise ValueError("Provide an https YouTube URL.")
    value = (
        parsed.path.strip("/") if parsed.hostname == "youtu.be" else parse_qs(parsed.query).get("v", [""])[0]
    )
    if not re.fullmatch(r"[\w-]{11}", value):
        raise ValueError("Invalid YouTube video ID.")
    return value


def tool_schema(name, description, properties, required):
    return {
        "type": "function",
        "function": {
            "name": name,
            "description": description,
            "parameters": {
                "type": "object",
                "properties": properties,
                "required": required,
                "additionalProperties": False,
            },
        },
    }


STRING = {"type": "string"}
INTEGER = {"type": "integer"}
SCHEMAS = [
    tool_schema(
        "web_search",
        "Search the public web for primary evidence. Do not search for benchmark answer keys or other submissions.",
        {"query": STRING},
        ["query"],
    ),
    tool_schema(
        "read_web",
        "Read a webpage/PDF and its source links. start is a character offset; use it to continue reading.",
        {"url": STRING, "start": INTEGER},
        ["url"],
    ),
    tool_schema(
        "read_file",
        "Read a registered attachment: PDF, XLSX, DOCX, PPTX, text, Python source, or CSV. Source code is returned as text. start is a character offset. For XLSX set formulas=true to inspect source formulas when cached values are missing.",
        {"filename": STRING, "start": INTEGER, "formulas": {"type": "boolean"}},
        ["filename"],
    ),
    tool_schema(
        "spreadsheet_sum",
        "Sum numeric cells exactly in specified rows of an XLSX column. Read the sheet first and select rows using question criteria. Returns included values and total.",
        {"filename": STRING, "sheet": STRING, "column": STRING, "rows": {"type": "array", "items": INTEGER}},
        ["filename", "sheet", "column", "rows"],
    ),
    tool_schema(
        "calculator",
        "Compute numeric arithmetic; supports + - * / // % **, lists, sum/min/max/round/sqrt/floor/ceil/abs/len.",
        {"expression": STRING},
        ["expression"],
    ),
    tool_schema(
        "text_operation",
        "Reverse text or alphabetically sort a comma-separated list.",
        {"text": STRING, "operation": {"type": "string", "enum": ["reverse", "sort_list"]}},
        ["text", "operation"],
    ),
    tool_schema(
        "inspect_image",
        "Ask a vision model about a registered image attachment or public image URL. For chess request all piece locations, side to move, and FEN; verify before using analyze_chess.",
        {"source": STRING, "instruction": STRING},
        ["source", "instruction"],
    ),
    tool_schema(
        "transcribe_audio",
        "Transcribe a registered audio attachment using the configured speech recognition provider.",
        {"filename": STRING},
        ["filename"],
    ),
    tool_schema(
        "youtube_transcript",
        "Read timestamped YouTube captions. If captions are unavailable use inspect_video with mode=audio to transcribe the audio.",
        {"url": STRING},
        ["url"],
    ),
    tool_schema(
        "inspect_video",
        "Inspect a short YouTube video: mode=frames samples up to 12 specified timestamps using vision; mode=audio transcribes the audio. Sampling can miss events, refine timestamps when necessary.",
        {
            "url": STRING,
            "mode": {"type": "string", "enum": ["frames", "audio"]},
            "instruction": STRING,
            "timestamps": {"type": "array", "items": {"type": "number"}},
        },
        ["url", "mode", "instruction"],
    ),
    tool_schema(
        "analyze_chess",
        "Find the best move from a verified FEN with Stockfish. Reports SAN, evaluation, and principal variation.",
        {"fen": STRING},
        ["fen"],
    ),
]


class Tools:
    def __init__(self, model, attachments: list[Path] | None = None):
        self.model = model
        self.files = {p.name: p for p in attachments or []}
        self.documents = {}

    def dispatch(self, name: str, arguments: dict):
        try:
            if name not in {s["function"]["name"] for s in SCHEMAS}:
                raise ValueError("Unknown tool.")
            result = getattr(self, name)(**arguments)
            return {"ok": True, "result": result}
        except InferenceAccountError:
            raise
        except Exception as exc:
            return {
                "ok": False,
                "error": safe_error(exc, self.model.settings.token),
                "hint": "Try another source or tool; do not assume missing evidence.",
            }

    def path(self, filename):
        if filename not in self.files:
            raise ValueError("Unknown attachment. Use the exact filename listed in the task.")
        return self.files[filename]

    def page(self, text, start=0):
        start = max(0, int(start))
        return {
            "text": text[start : start + 16000],
            "start": start,
            "total_characters": len(text),
            "next_start": start + 16000 if start + 16000 < len(text) else None,
        }

    def web_search(self, query):
        if not 1 <= len(query) <= 1000:
            raise ValueError("Search query must be 1..1000 characters.")
        rows = list(DDGS(timeout=20).text(query, max_results=6))
        if not rows:
            raise ValueError("Search returned no results; rephrase the query.")
        return [
            {"title": r.get("title"), "url": r.get("href"), "snippet": r.get("body", "")[:2000]} for r in rows
        ]

    def read_web(self, url, start=0):
        if url not in self.documents:
            data, mime, final_url = fetch_url(url)
            if "pdf" in mime or data.startswith(b"%PDF"):
                text = extract_text(data, ".pdf")
            else:
                soup = BeautifulSoup(data, "html.parser")
                for element in soup(["script", "style", "nav", "footer", "header"]):
                    element.decompose()
                links = [
                    {"text": a.get_text(" ", strip=True)[:120], "url": urljoin(final_url, a["href"])}
                    for a in soup.select("a[href]")
                    if a["href"] and not a["href"].startswith(("#", "javascript:", "mailto:"))
                ]
                text = (
                    soup.get_text("\n", strip=True)
                    + "\n\nSOURCE LINKS:\n"
                    + json.dumps(links[:300], ensure_ascii=False)
                )
            self.documents[url] = f"SOURCE: {final_url}\n" + text[:600000]
        return self.page(self.documents[url], start)

    def read_file(self, filename, start=0, formulas=False):
        path = self.path(filename)
        cache_key = (filename, bool(formulas))
        if cache_key not in self.documents:
            self.documents[cache_key] = extract_text(
                path.read_bytes(), path.suffix.lower(), formulas=formulas
            )[:600000]
        return self.page(self.documents[cache_key], start)

    def spreadsheet_sum(self, filename, sheet, column, rows):
        if (
            not re.fullmatch(r"[A-Z]{1,3}", column)
            or not rows
            or len(rows) > 2000
            or len(set(rows)) != len(rows)
        ):
            raise ValueError("Provide a column letter and 1..2000 unique row numbers.")
        if any(type(r) is not int or not 1 <= r <= 100000 for r in rows):
            raise ValueError("Invalid spreadsheet row numbers.")
        data = self.path(filename).read_bytes()
        check_archive(data)
        book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
        try:
            values = []
            for row in sorted(rows):
                address = f"{column}{row}"
                value = book[sheet][address].value
                if type(value) not in {int, float}:
                    raise ValueError(f"{address} is not numeric or has no cached formula value.")
                values.append({"cell": address, "value": str(value)})
            total = sum((Decimal(v["value"]) for v in values), Decimal(0))
            return {"cells": values, "total": str(total), "two_decimals": f"{total:.2f}"}
        finally:
            book.close()

    def calculator(self, expression):
        return calculator(expression)

    def text_operation(self, text, operation):
        if len(text) > 20000:
            raise ValueError("Text is too long.")
        if operation == "reverse":
            return text[::-1]
        if operation == "sort_list":
            return ", ".join(sorted((s.strip() for s in text.split(",") if s.strip()), key=str.casefold))
        raise ValueError("Unsupported text operation.")

    def inspect_image(self, source, instruction):
        data = self.path(source).read_bytes() if source in self.files else fetch_url(source)[0]
        return self.model.inspect_images([data], instruction)

    def asr(self, data, filename="recording.mp3"):
        return self.model.transcribe_audio(data, filename)

    def transcribe_audio(self, filename):
        return self.asr(self.path(filename).read_bytes(), filename)

    def youtube_transcript(self, url):
        from youtube_transcript_api import YouTubeTranscriptApi

        video_id = youtube_id(url)
        api = YouTubeTranscriptApi()
        try:
            result = api.fetch(video_id, languages=["en", "en-US", "en-GB"])
        except Exception:
            result = next(iter(api.list(video_id))).fetch()
        return self.page("\n".join(f"[{s.start:.1f}s] {s.text}" for s in result))

    def inspect_video(self, url, mode, instruction, timestamps=None):
        video_id = youtube_id(url)
        if mode not in {"frames", "audio"}:
            raise ValueError("mode must be frames or audio.")
        timestamps = timestamps if timestamps is not None else [0, 15, 30, 45, 60, 90]
        if not 1 <= len(timestamps) <= 12 or any(
            not isinstance(t, (int, float)) or not 0 <= t <= 600 for t in timestamps
        ):
            raise ValueError("Use 1..12 timestamps, each between 0 and 600 seconds.")
        # yt-dlp runs in a separate process so the entire download has a hard timeout.
        import sys

        import imageio_ffmpeg

        with tempfile.TemporaryDirectory(prefix="course-video-") as folder:
            template = str(Path(folder) / "video.%(ext)s")
            subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "yt_dlp",
                    "--no-playlist",
                    "--quiet",
                    "--no-warnings",
                    "--socket-timeout",
                    "20",
                    "--retries",
                    "1",
                    "--max-filesize",
                    "75M",
                    "--match-filter",
                    "duration <= 600",
                    "--no-cache-dir",
                    "--no-write-info-json",
                    "-f",
                    "best[height<=480][ext=mp4]/best[height<=480]",
                    "-o",
                    template,
                    f"https://www.youtube.com/watch?v={video_id}",
                ],
                check=True,
                capture_output=True,
                timeout=120,
            )
            videos = [p for p in Path(folder).glob("video.*") if p.suffix not in {".part", ".ytdl"}]
            if not videos:
                raise ValueError("Video is unavailable, too large, or longer than 10 minutes.")
            ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
            if mode == "audio":
                audio = Path(folder) / "audio.wav"
                subprocess.run(
                    [ffmpeg, "-y", "-i", str(videos[0]), "-vn", "-ac", "1", "-ar", "16000", str(audio)],
                    check=True,
                    capture_output=True,
                    timeout=60,
                )
                if audio.stat().st_size > 25 * 1024 * 1024:
                    raise ValueError("Audio exceeds the transcription size limit.")
                return self.asr(audio.read_bytes(), audio.name)
            frames, sampled = [], []
            for i, timestamp in enumerate(timestamps):
                target = Path(folder) / f"frame-{i}.jpg"
                subprocess.run(
                    [
                        ffmpeg,
                        "-y",
                        "-ss",
                        str(timestamp),
                        "-i",
                        str(videos[0]),
                        "-frames:v",
                        "1",
                        "-vf",
                        "scale=960:-2",
                        str(target),
                    ],
                    check=True,
                    capture_output=True,
                    timeout=20,
                )
                if target.exists():
                    frames.append(target.read_bytes())
                    sampled.append(timestamp)
            if not frames:
                raise ValueError("No frames at the requested timestamps.")
            observation = self.model.inspect_images(
                frames,
                f"{instruction}\nFrame timestamps in seconds, in order: {sampled}. These are samples and may miss intervening events.",
            )
            return {
                "timestamps": sampled,
                "observations": observation,
                "limitation": "Sampled frames do not prove an event never occurs between samples.",
            }

    def analyze_chess(self, fen):
        import chess
        import chess.engine

        engine_path = os.getenv("STOCKFISH_PATH") or shutil.which("stockfish")
        if not engine_path and Path("/usr/games/stockfish").exists():
            engine_path = "/usr/games/stockfish"
        if not engine_path:
            raise ValueError(
                "Stockfish is unavailable. Set STOCKFISH_PATH locally; Spaces installs it via packages.txt."
            )
        board = chess.Board(fen)
        if not board.is_valid() or board.is_game_over():
            raise ValueError("FEN is invalid or the game has already ended.")
        with chess.engine.SimpleEngine.popen_uci(engine_path, timeout=15) as engine:
            info = engine.analyse(board, chess.engine.Limit(time=5))
            variation = info["pv"][:10]
            return {
                "best_move_san": board.san(variation[0]),
                "score_side_to_move": str(info["score"].pov(board.turn)),
                "variation": board.variation_san(variation),
                "fen": fen,
            }
